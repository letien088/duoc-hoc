"""Chuyển 2 file Excel của khoá học thành data/khoahoc.json cho tab 🎓 Khoá học.

Chạy lại mỗi khi Excel đổi:   python tools/xuat_khoahoc.py
rồi đổi số BAN trong sw.js và đẩy lên như README mục 4.

Excel có 16 sheet. Sheet nào đi vào app, sheet nào không, và vì sao:
  Data HV final ............ Bộ HV (bộ chính của học viên, 34 nhóm)
  Lọc data học viên ........ tuần học (Tuần chạy data 1-3) của từng nhóm trong bộ HV
  Data 700 HV / 1500 / 3000  ba bộ còn lại (Level 1-2-3)
  Bổ sung tt Data bán chạy . đánh dấu 🔥 bán chạy + Brandname/Generic theo mã SP
  Data Phương gửi .......... tra tên / nước sản xuất cho mã SP (Data 1500 không ghi tên)
  Nhóm & dạng dùng ......... thứ tự nhóm chuẩn (đúng thứ tự dạy)
  AI cắt liều .............. tình huống cắt liều (đề + đáp án + từ khoá) và chủ đề
  Tổng lộ trình ............ cấu trúc khoá học + các tình huống QTBH
  Phân chia team chuyên môn  lịch bài giảng và thời lượng (BỎ tên giảng viên)
  KHÔNG đưa vào: Data bỏ (dòng đã loại khỏi bộ HV), Timeline, Mẫu lesson plan,
  Hướng dẫn làm Data, Bảng tổng hợp 3 — đây là giấy tờ làm việc của nhóm soạn,
  không phải nội dung học. Tên giảng viên cũng không đưa vào app (app công khai).
"""
import datetime
import json
import os
import re
import sys
import unicodedata

import openpyxl

# ============================================================================
# |                          BẢNG ĐIỀU KHIỂN                                 |
# ============================================================================
DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # thư mục app

# File Excel đầy đủ (16 sheet). VD: D:\congviec\code\duoc_hoc\6.4.2026 Bộ data 700 - 1500 - 3000.xlsx
FILE_DAY_DU = os.path.join(DIR, '6.4.2026 Bộ data 700 - 1500 - 3000.xlsx')
# File học viên (14 sheet, trùng với file trên) — chỉ dùng khi file đầy đủ thiếu sheet
# VD: D:\congviec\code\duoc_hoc\6.4.2026 data học viên.xlsx
FILE_HOC_VIEN = os.path.join(DIR, '6.4.2026 data học viên.xlsx')
# File danh mục Long Châu đã có trong app — để tra tên / nước khi Excel để trống
FILE_DANH_MUC = os.path.join(DIR, 'data', 'danhmuc.json')
# Nơi ghi kết quả. VD: D:\congviec\code\duoc_hoc\data\khoahoc.json
FILE_RA = os.path.join(DIR, 'data', 'khoahoc.json')
NGAY_DU_LIEU = '06/04/2026'        # ngày ghi trên tên file Excel

# Bốn bộ, theo thứ tự hiện trong app: (khoá, sheet, tên, mô tả)
CAC_BO = [
    ('hv',   'Data HV final', 'Bộ HV',
     'Bộ chính thức của học viên: 34 nhóm, chia theo 3 tuần như lịch khoá học.'),
    ('700',  'Data 700 HV',   'Bộ 700',
     'Level 1 — khoảng 700 biệt dược phổ biến nhất, 26 nhóm.'),
    ('1500', 'Data 1500',     'Bộ 1500',
     'Level 2 — mở rộng lên ~1.500 biệt dược, thêm thần kinh, gan mật, bổ, nhỏ mắt…'),
    ('3000', 'Data 3000',     'Bộ 3000',
     'Level 3 — ~3.000 biệt dược kèm sản phẩm đi kèm (tuýp bôi, kẹo ngậm, sản phẩm hỗ trợ).'),
]
# ============================================================================


def dat_tieu_de():
    try:
        import ctypes
        ctypes.windll.kernel32.SetConsoleTitleW('tools/xuat_khoahoc')
    except Exception:
        pass


def sach(v):
    """Chuỗi gọn: bỏ xuống dòng thừa, khoảng trắng đôi. None/rỗng/lỗi Excel -> ''."""
    if v is None:
        return ''
    s = str(v).replace('\xa0', ' ')
    s = re.sub(r'\s+', ' ', s).strip()
    if s in ('#N/A', '#REF!', '#VALUE!', 'None', '-'):
        return ''
    return s


def bo_dau(s):
    s = unicodedata.normalize('NFD', s.lower()).replace('đ', 'd')
    return ''.join(c for c in s if unicodedata.category(c) != 'Mn')


def ham_luong(cell):
    """Ô phần trăm Excel lưu số 0.0005 kèm định dạng '0.00%' -> '0.05%'."""
    v = cell.value
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if '%' in (cell.number_format or ''):
            return f'{v * 100:.4g}%'
        return f'{v:g}'
    return sach(v)


def ma_sp(v):
    """Mã sản phẩm Long Châu: số -> đủ 8 chữ số như SKU trong danh mục."""
    s = sach(v).split(' ')[0]
    s = re.sub(r'\.0$', '', s)
    if not re.fullmatch(r'[0-9A-Za-z]+', s) or set(s) <= {'0'}:
        return ''
    return s.zfill(8) if s.isdigit() else s.upper()


def tach_tu_khoa(s):
    """Tách từ khoá theo dấu phẩy / xuống dòng, nhưng KHÔNG tách phần trong ngoặc:
    "Theo dõi (sốt, phân lẫn nhầy, máu) nếu không đỡ đi khám" là một từ khoá."""
    ra, cur, sau = [], '', 0
    for ch in s:
        if ch == '(':
            sau += 1
        elif ch == ')':
            sau = max(0, sau - 1)
        if ch in ',\n' and not sau:
            ra.append(cur)
            cur = ''
        else:
            cur += ch
    ra.append(cur)
    return ra


def hoa_dau(s):
    return s[:1].upper() + s[1:] if s else s


class Doc:
    """Đọc sheet theo tên, lấy ở file đầy đủ trước, thiếu thì lấy ở file học viên."""

    def __init__(self):
        self.wb = []
        for f in (FILE_DAY_DU, FILE_HOC_VIEN):
            if os.path.exists(f):
                self.wb.append(openpyxl.load_workbook(f, read_only=True))
        if not self.wb:
            sys.exit('Không thấy file Excel nào — kiểm tra FILE_DAY_DU / FILE_HOC_VIEN.')

    def sheet(self, ten):
        for wb in self.wb:
            if ten in wb.sheetnames:
                return [list(r) for r in wb[ten].iter_rows()]
        sys.exit(f'Không thấy sheet "{ten}" trong file Excel nào.')


def main():
    dat_tieu_de()
    try:
        sys.stdout.reconfigure(encoding='utf-8')   # cmd mặc định cp1252 không in được tiếng Việt
    except Exception:
        pass
    doc = Doc()

    # --- Nguồn tra cứu phụ theo mã SP -----------------------------------------
    dm = {}
    if os.path.exists(FILE_DANH_MUC):
        g = json.load(open(FILE_DANH_MUC, encoding='utf-8'))
        c = g['cot']
        for t in g['t']:
            dm[t[c.index('sku')]] = (t[c.index('ten')], t[c.index('xuatXu')], t[c.index('keDon')])
    phuong = {}
    for r in doc.sheet('Data Phương gửi')[1:]:
        m = ma_sp(r[0].value)
        if m:
            phuong[m] = (sach(r[2].value), sach(r[6].value))

    # --- Sản phẩm (biệt dược): mỗi mã SP một bản ghi dùng chung cho mọi bộ ----
    sp, sp_theo_ma = [], {}

    def them_sp(ma, ten, nuoc='', loai=''):
        ten = sach(ten)
        nuoc = sach(nuoc)
        if nuoc.lower() in ('x', '') or re.fullmatch(r'[\d.]+', nuoc):
            nuoc = ''
        loai = bo_dau(sach(loai)).replace(' ', '')
        loai = 1 if loai == 'brandname' else 0 if loai == 'generic' else None
        k = ma or ('ten:' + bo_dau(ten))
        if not ma and not ten:
            return None
        if k not in sp_theo_ma:
            sp_theo_ma[k] = len(sp)
            sp.append({'ma': ma, 'ten': ten, 'nuoc': nuoc, 'loai': loai, 'banChay': 0})
        x = sp[sp_theo_ma[k]]
        if ten and not x['ten']:
            x['ten'] = ten
        if nuoc and not x['nuoc']:
            x['nuoc'] = nuoc
        if loai is not None and x['loai'] is None:
            x['loai'] = loai
        return sp_theo_ma[k]

    def doc_bd(tieu_de, r):
        """Các biệt dược trên một dòng: cột 'Biệt dược…/BD…' mở biệt dược mới, 'MSP' là
        mã của nó, 'Nước sản xuất' và 'Phân loại' gắn vào biệt dược gần nhất bên trái."""
        ds, cur = [], None
        for j, h in enumerate(tieu_de):
            h = bo_dau(sach(h))
            if j >= len(r):
                break
            v = r[j].value
            # "Biệt dược 1", "BD tương đương 2", "Biêt dược theo toa" (gõ sai dấu) — bỏ dấu là như nhau
            if h.startswith('biet duoc') or h.startswith('bd '):
                cur = {'ten': sach(v), 'ma': '', 'nuoc': '', 'loai': ''}
                ds.append(cur)
            elif h.startswith('msp'):
                if cur is None or cur['ma']:
                    cur = {'ten': '', 'ma': '', 'nuoc': '', 'loai': ''}
                    ds.append(cur)
                cur['ma'] = ma_sp(v)
            elif h.startswith('nuoc san xuat') and cur is not None:
                cur['nuoc'] = sach(v)
            elif h.startswith('phan loai') and cur is not None:
                cur['loai'] = sach(v)
        ra = []
        for b in ds:
            i = them_sp(b['ma'], b['ten'], b['nuoc'], b['loai'])
            if i is not None and i not in ra:
                ra.append(i)
        return ra

    # Brandname/Generic và "bán chạy" nằm ở các sheet làm việc: đọc trước để gắn theo mã
    loc = doc.sheet('Lọc data học viên')
    tuan_nhom = {}
    for r in loc[1:]:
        if len(r) < 17 or not sach(r[1].value):
            continue
        # cột: Biêt dược theo toa(8) · MSP(10) · BD tương đương 1(11) · MSP 1(12) · Nước(13) · Phân loại(14)
        them_sp(ma_sp(r[10].value), r[8].value)
        them_sp(ma_sp(r[12].value), r[11].value, r[13].value, r[14].value)
        t = sach(r[4].value)
        if t.isdigit():
            tuan_nhom.setdefault(sach(r[1].value), int(t))
    ban_chay = doc.sheet('Bổ sung tt Data bán chạy')
    for r in ban_chay[1:]:
        if len(r) < 9:
            continue
        i = them_sp(ma_sp(r[6].value), r[5].value, r[7].value, r[8].value)
        if i is not None:
            sp[i]['banChay'] = 1

    # --- Nhóm: thứ tự chuẩn lấy từ sheet "Nhóm & dạng dùng" -----------------
    nd = doc.sheet('Nhóm & dạng dùng')
    thu_tu_nhom = []
    for j in (0, 1):
        for r in nd[1:]:
            s = sach(r[j].value)
            if s and s not in thu_tu_nhom:
                thu_tu_nhom.append(s)
    DOI_TEN = {'Tiền liệt tuyến': 'Tuyến tiền liệt', 'Nsaids': 'NSAIDs'}
    thu_tu_nhom = [DOI_TEN.get(n, n) for n in thu_tu_nhom]
    # Bộ HV gộp "Ho đàm"/"Ho khan" thành "Ho"; nhóm chưa có trong danh sách thì xếp cuối
    if 'Ho' not in thu_tu_nhom:
        thu_tu_nhom.insert(thu_tu_nhom.index('Ho đàm') if 'Ho đàm' in thu_tu_nhom else 0, 'Ho')
    tuan_nhom = {DOI_TEN.get(k, k): v for k, v in tuan_nhom.items()}
    for k in ('Ho đàm', 'Ho khan'):
        if k in tuan_nhom:
            tuan_nhom.setdefault('Ho', tuan_nhom[k])

    nhom, pn = [], []

    def so_nhom(ten):
        if ten not in nhom:
            nhom.append(ten)
        return nhom.index(ten)

    def so_pn(ten):
        if ten not in pn:
            pn.append(ten)
        return pn.index(ten)

    # --- Bốn bộ ---------------------------------------------------------------
    bo_ra = []
    muc = []                  # mọi dòng của mọi bộ: [nhóm, phân nhóm, dược chất, hàm lượng, dạng, [sp]]
    for khoa, ten_sheet, ten, mo in CAC_BO:
        rows = doc.sheet(ten_sheet)
        td = [c.value for c in rows[0]]
        cac = []
        for r in rows[1:]:
            if len(r) < 7:
                continue
            n = DOI_TEN.get(sach(r[1].value), sach(r[1].value))
            if not n:
                continue
            p = DOI_TEN.get(sach(r[2].value), sach(r[2].value)) or n
            dc = sach(r[3].value)
            hl = ham_luong(r[4])
            dang = hoa_dau(sach(r[5].value))
            bd = doc_bd(td, r)
            if not dc and not bd:
                continue
            cac.append((n, p, dc, hl, dang, bd))

        def khoa_sap(x, cac=cac):
            n = x[1][0]
            return (tuan_nhom.get(n, 9) if khoa == 'hv' else 0,
                    thu_tu_nhom.index(n) if n in thu_tu_nhom else 999, x[0])
        cac = [c for _, c in sorted(enumerate(cac), key=khoa_sap)]
        ids = []
        for n, p, dc, hl, dang, bd in cac:
            ids.append(len(muc))
            muc.append([so_nhom(n), so_pn(p), dc, hl, dang, bd])
        bo_ra.append({'k': khoa, 'ten': ten, 'mo': mo, 'muc': ids})

    # Chỉ giữ biệt dược mà ít nhất một bộ dùng tới. Các sheet làm việc (lọc data, bán chạy)
    # chỉ để gắn thông tin theo mã, không được thêm thuốc lạ — vd dòng đã bị đề xuất bỏ.
    dung = sorted({i for m in muc for i in m[5]})
    so_moi = {cu: moi for moi, cu in enumerate(dung)}
    sp = [sp[i] for i in dung]
    for m in muc:
        m[5] = [so_moi[i] for i in m[5]]

    # Tên / nước còn trống: lấy từ bảng mã SP của Phương, rồi từ danh mục Long Châu
    for x in sp:
        if x['ma'] in phuong:
            t, nuoc = phuong[x['ma']]
            x['ten'] = x['ten'] or t
            x['nuoc'] = x['nuoc'] or nuoc
        if x['ma'] in dm:
            t, nuoc, _ = dm[x['ma']]
            x['ten'] = x['ten'] or sach(t)
            x['nuoc'] = x['nuoc'] or sach(nuoc)
        x['ten'] = x['ten'] or ('Mã SP ' + x['ma'])

    # --- Tình huống cắt liều + chủ đề ----------------------------------------
    cl = doc.sheet('AI cắt liều')
    chu_de, tinh_huong, dt = [], [], ''
    cach_lam = ''
    bat_dau_th = None
    for i, r in enumerate(cl):
        v = [sach(c.value) for c in r[:9]] + [''] * 9
        if v[0] == 'Loại' and v[1] == 'Câu hỏi':
            bat_dau_th = i + 1
            break
        if v[0] == 'Cách làm tình huống':
            nxt = cl[i + 1][0].value if i + 1 < len(cl) else ''
            cach_lam = str(nxt or '').strip()
            continue
        if i and v[1] and not v[0].startswith('Đưa ra'):
            dt = v[0] or dt
            chu_de.append({'dt': dt, 'benh': v[1], 'ghiChu': v[2]})
    if bat_dau_th:
        for r in cl[bat_dau_th:]:
            v = [str(c.value).strip() if c.value is not None else '' for c in r[:9]] + [''] * 9
            if not v[1]:
                continue
            dong_dap = [d.strip() for d in v[2].split('\n') if d.strip()]
            hoi = [d.strip() for d in v[1].split('\n') if d.strip()]
            # hai câu hỏi cố định ở cuối đề tách riêng cho gọn
            cau = [d.lstrip('- ').strip() for d in hoi if d.lstrip('- ').startswith('Bạn')]
            de = ' '.join(d for d in hoi if not d.lstrip('- ').startswith('Bạn'))
            tk = [sach(t) for t in tach_tu_khoa(v[3]) if sach(t)]
            tinh_huong.append({
                'loai': sach(v[0]), 'de': de, 'cau': cau,
                'benh': sach(dong_dap[0]) if dong_dap else '', 'loiKhuyen': [sach(d) for d in dong_dap[1:]],
                'tuKhoa': tk, 'level': sach(v[4]), 'doiTuong': sach(v[5]),
            })

    # --- Tổng lộ trình ---------------------------------------------------------
    tl = doc.sheet('Tổng lộ trình')
    h = [sach(c.value) for c in tl[0]]
    hv = [str(c.value).strip() if c.value is not None else '' for c in tl[1]]
    def gon_dong(s):
        # "34 nhóm: \n- 250 HC + Khác 69 \n- Tổng: 700 BD" -> "34 nhóm · 250 HC + Khác 69 · Tổng: 700 BD"
        ds = [re.sub(r'^[-+]\s*', '', d).strip().rstrip(':').strip() for d in str(s).split('\n')]
        return ' · '.join(d for d in ds if d)

    tong_quan = []
    for j in range(1, 12, 2):
        if j < len(h) and h[j] and hv[j]:
            tong_quan.append([h[j], gon_dong(hv[j])])
    qtbh = []
    for j, t in enumerate(h):
        if t == 'Tình huống QTBH' and j + 1 < len(hv):
            for d in hv[j + 1].split('\n'):
                d = d.strip()
                if d.startswith('+'):
                    qtbh.append(d.lstrip('+ ').strip())
    cap = []
    for r in tl[1:]:
        a, d = sach(r[0].value), (str(r[11].value).strip() if len(r) > 11 and r[11].value else '')
        if a and d:
            cap.append(['Học viên' if a == 'HV' else a, gon_dong(d)])

    # --- Lịch bài giảng (bỏ tên giảng viên) ------------------------------------
    pc = doc.sheet('Phân chia team chuyên môn')
    lich, dot = [], None
    for r in pc[2:]:
        a = str(r[0].value).strip() if r[0].value else ''
        b = sach(r[1].value) if len(r) > 1 else ''
        if a:
            ten_dot = ' · '.join(
                re.sub(r'^-\s*', '', d).strip() for d in a.split('\n')
                if d.strip() and not re.fullmatch(r'\(.*quản\)', d.strip()))
            ten_dot = re.sub(r'\s*\(.*?quản\)', '', ten_dot).strip(' ·')
            ten_dot = (ten_dot.replace('Tài liệu VX &', 'Vaccine').replace('Tuần 3; 4', 'Tuần 3–4')
                       .replace('Tuần 4; 5', 'Tuần 4–5'))
            dot = {'ten': ten_dot or 'Vận hành nhà thuốc', 'bai': []}
            lich.append(dot)
            if a.startswith('Nội dung tuần 1'):
                # ô bên cạnh là tên giảng viên, không phải bài; số bài lấy ở sheet Tổng lộ trình
                so = next((v for k, v in tong_quan if k == 'Bài giảng tuần 1'), '')
                dot['bai'].append([f'{so} tổng quan' if so else 'Bài giảng tổng quan', ''])
                continue
        if dot is not None and b:
            m = re.match(r'^(.*?)[\s:-]*((?:\d+\s*h\s*\d*\s*(?:phút|ph)?)|(?:\d+\s*(?:phút|ph)))$', b)
            if m:
                tg = re.sub(r'\s*(phút|ph)$', ' phút', m.group(2).strip())
                dot['bai'].append([m.group(1).strip(), tg])
            else:
                dot['bai'].append([b, ''])

    ra = {
        'v': 1, 'ngay': NGAY_DU_LIEU,
        'nguon': 'Bộ data 700 - 1500 - 3000 (Excel 6.4.2026)',
        'taoLuc': datetime.date.today().strftime('%d/%m/%Y'),
        'nhom': nhom, 'pn': pn,
        'tuan': {str(nhom.index(k)): v for k, v in tuan_nhom.items() if k in nhom},
        # [mã, tên, nước, 1 Brandname / 0 Generic / null, 1 bán chạy, 1 có trong danh mục, 1 kê đơn / 0 không / null]
        'sp': [[x['ma'], x['ten'], x['nuoc'], x['loai'], x['banChay'], 1 if x['ma'] in dm else 0,
                dm[x['ma']][2] if x['ma'] in dm else None] for x in sp],
        'muc': muc, 'bo': bo_ra,
        'tinhHuong': tinh_huong, 'chuDe': chu_de, 'cachLam': cach_lam,
        'tongQuan': tong_quan, 'cap': cap, 'qtbh': qtbh, 'lich': lich,
    }
    with open(FILE_RA, 'w', encoding='utf-8') as f:
        json.dump(ra, f, ensure_ascii=False, separators=(',', ':'))

    print(f'Đã ghi {FILE_RA} ({os.path.getsize(FILE_RA) / 1024:.0f} KB)')
    for b in bo_ra:
        ds = [muc[i] for i in b['muc']]
        print(f"  {b['ten']:8} {len(ds):5} dòng · {len({m[0] for m in ds}):2} nhóm · "
              f"{len({m[2].lower() for m in ds if m[2]}):4} dược chất · "
              f"{len({s for m in ds for s in m[5]}):5} biệt dược")
    print(f'  {len(sp)} biệt dược (mã SP) · {sum(1 for x in sp if x["ma"] in dm)} có trong danh mục Long Châu · '
          f'{sum(1 for x in sp if x["loai"] is not None)} biết Brandname/Generic · {sum(x["banChay"] for x in sp)} bán chạy')
    print(f'  {len(tinh_huong)} tình huống cắt liều · {len(chu_de)} chủ đề · {len(qtbh)} tình huống QTBH · '
          f'{sum(len(d["bai"]) for d in lich)} bài giảng trong {len(lich)} đợt')


if __name__ == '__main__':
    main()
